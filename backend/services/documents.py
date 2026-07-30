"""Safe in-memory text extraction for documents uploaded for analysis."""
from __future__ import annotations

import io
import re
from dataclasses import dataclass, field
from pathlib import Path
from uuid import uuid4

from ..core.config import settings
from ..core.token_usage import count_text

TEXT_EXTENSIONS = {
    ".txt", ".md", ".csv", ".json", ".xml", ".yaml", ".yml", ".html", ".htm",
    ".rtf", ".log", ".py", ".js", ".ts", ".java", ".c", ".cpp", ".sql",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {".pdf", ".docx", ".pptx"}
LEGACY_EXTENSIONS = {".doc", ".ppt"}


class DocumentExtractionError(ValueError):
    """A user-correctable invalid or unreadable document."""


@dataclass
class ExtractedDocument:
    filename: str
    extension: str
    content_type: str
    size_bytes: int
    text: str
    pages_or_slides: int | None = None
    truncated: bool = False
    warnings: list[str] = field(default_factory=list)
    saved_path: str | None = None

    def public(self) -> dict:
        return {
            "filename": self.filename,
            "extension": self.extension,
            "content_type": self.content_type,
            "size_bytes": self.size_bytes,
            "text": self.text,
            "characters": len(self.text),
            "estimated_tokens": count_text(self.text),
            "pages_or_slides": self.pages_or_slides,
            "truncated": self.truncated,
            "warnings": self.warnings,
            "saved_path": self.saved_path,
        }


def save_upload(
    session_id: str,
    filename: str,
    content: bytes,
) -> str:
    """Persist an already validated chat upload below its session directory."""
    safe_session = re.sub(r"[^a-zA-Z0-9_-]", "", session_id)[:64]
    if not safe_session:
        raise DocumentExtractionError("Invalid session id for upload storage.")
    safe_name = Path(filename or "document").name
    target_dir = Path(settings.chat_upload_dir) / safe_session
    target_dir.mkdir(parents=True, exist_ok=True)
    stored_name = f"{uuid4().hex[:12]}-{safe_name}"
    path = target_dir / stored_name
    path.write_bytes(content)
    return str(path.relative_to(Path(settings.chat_upload_dir).parent))


def extract(filename: str, content: bytes, content_type: str) -> ExtractedDocument:
    safe_name = Path(filename or "document").name
    extension = Path(safe_name).suffix.lower()
    if not content:
        raise DocumentExtractionError("The uploaded document is empty.")
    if len(content) > settings.max_upload_bytes:
        limit_mb = settings.max_upload_bytes // (1024 * 1024)
        raise DocumentExtractionError(f"The document exceeds the {limit_mb} MB limit.")
    if extension in LEGACY_EXTENSIONS:
        modern = ".docx" if extension == ".doc" else ".pptx"
        raise DocumentExtractionError(
            f"Legacy {extension} files are not parsed safely. Save the file as {modern} "
            "and upload it again."
        )
    if extension not in SUPPORTED_EXTENSIONS:
        supported = ", ".join(sorted(SUPPORTED_EXTENSIONS))
        raise DocumentExtractionError(
            f"Unsupported file type '{extension or 'none'}'. Supported: {supported}."
        )

    warnings: list[str] = []
    pages_or_slides = None
    try:
        if extension == ".pdf":
            text, pages_or_slides = _pdf(content)
            if not re.sub(r"--- Page \d+ ---", "", text).strip():
                warnings.append(
                    "No selectable text was found. Scanned PDFs require OCR, which is not enabled."
                )
        elif extension == ".docx":
            text = _docx(content)
        elif extension == ".pptx":
            text, pages_or_slides = _pptx(content)
        else:
            text = _text(content, extension)
    except DocumentExtractionError:
        raise
    except Exception as error:
        raise DocumentExtractionError(
            f"Could not read '{safe_name}': {type(error).__name__}: {error}"
        ) from error

    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    truncated = len(text) > settings.max_document_chars
    if truncated:
        text = text[: settings.max_document_chars]
        warnings.append(
            f"Extracted text was truncated to {settings.max_document_chars:,} characters."
        )
    return ExtractedDocument(
        filename=safe_name,
        extension=extension,
        content_type=content_type or "application/octet-stream",
        size_bytes=len(content),
        text=text,
        pages_or_slides=pages_or_slides,
        truncated=truncated,
        warnings=warnings,
    )


def _pdf(content: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    pages = [
        f"--- Page {index} ---\n{page.extract_text() or ''}"
        for index, page in enumerate(reader.pages, start=1)
    ]
    return "\n\n".join(pages), len(reader.pages)


def _docx(content: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(content))
    lines = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return "\n".join(lines)


def _pptx(content: bytes) -> tuple[str, int]:
    from pptx import Presentation

    presentation = Presentation(io.BytesIO(content))
    slides: list[str] = []
    for index, slide in enumerate(presentation.slides, start=1):
        blocks: list[str] = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                value = shape.text.strip()
                if value:
                    blocks.append(value)
            if getattr(shape, "has_table", False):
                for row in shape.table.rows:
                    blocks.append(" | ".join(cell.text.strip() for cell in row.cells))
        notes = getattr(slide, "notes_slide", None)
        if notes is not None:
            note_text = [
                shape.text.strip()
                for shape in notes.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            if note_text:
                blocks.append("Speaker notes:\n" + "\n".join(note_text))
        slides.append(f"--- Slide {index} ---\n" + "\n".join(blocks))
    return "\n\n".join(slides), len(presentation.slides)


def _text(content: bytes, extension: str) -> str:
    text = content.decode("utf-8-sig", errors="replace")
    if extension in {".html", ".htm"}:
        from ..services.web import _TextExtractor

        parser = _TextExtractor()
        parser.feed(text)
        return "\n".join(parser.chunks)
    if extension == ".rtf":
        text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", text)
        text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
        text = text.replace("{", "").replace("}", "")
    return text
