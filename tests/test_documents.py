from __future__ import annotations

import io
import unittest

from docx import Document
from pptx import Presentation
from pptx.util import Inches
from pypdf import PdfWriter

from backend.services.documents import (
    DocumentExtractionError,
    extract,
)


class DocumentExtractionTests(unittest.TestCase):
    def test_extracts_utf8_text(self) -> None:
        result = extract("notes.txt", "Salut, document!".encode(), "text/plain")

        self.assertEqual(result.text, "Salut, document!")
        self.assertGreater(result.public()["estimated_tokens"], 0)

    def test_extracts_docx_paragraphs_and_tables(self) -> None:
        document = Document()
        document.add_heading("Candidate", 1)
        document.add_paragraph("Strong Python experience")
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Skill"
        table.cell(0, 1).text = "Level"
        payload = io.BytesIO()
        document.save(payload)

        result = extract(
            "candidate.docx",
            payload.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )

        self.assertIn("Strong Python experience", result.text)
        self.assertIn("Skill | Level", result.text)

    def test_extracts_pptx_slide_text(self) -> None:
        presentation = Presentation()
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        textbox = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1))
        textbox.text = "Architecture review"
        payload = io.BytesIO()
        presentation.save(payload)

        result = extract(
            "review.pptx",
            payload.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )

        self.assertEqual(result.pages_or_slides, 1)
        self.assertIn("Architecture review", result.text)

    def test_scanned_or_blank_pdf_returns_warning(self) -> None:
        writer = PdfWriter()
        writer.add_blank_page(width=200, height=200)
        payload = io.BytesIO()
        writer.write(payload)

        result = extract("blank.pdf", payload.getvalue(), "application/pdf")

        self.assertEqual(result.pages_or_slides, 1)
        self.assertTrue(result.warnings)

    def test_legacy_office_format_has_conversion_guidance(self) -> None:
        with self.assertRaisesRegex(DocumentExtractionError, r"\.docx"):
            extract("old.doc", b"legacy", "application/msword")


if __name__ == "__main__":
    unittest.main()
