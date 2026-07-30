"""Generate downloadable answer documents for the admin chat."""
from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from xml.sax.saxutils import escape


@dataclass(frozen=True)
class GeneratedDocument:
    content: bytes
    media_type: str
    filename: str


def generate(kind: str, answer: str, metadata: dict | None = None) -> GeneratedDocument:
    kind = kind.lower()
    metadata = metadata or {}
    if kind == "pdf":
        content = _pdf(answer, metadata)
        media_type = "application/pdf"
    elif kind == "docx":
        content = _docx(answer, metadata)
        media_type = (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        )
    elif kind == "json":
        content = json.dumps(
            {"answer": answer, "metadata": metadata},
            ensure_ascii=False,
            indent=2,
        ).encode("utf-8")
        media_type = "application/json"
    elif kind == "md":
        content = _text_document(answer, metadata, markdown=True).encode("utf-8")
        media_type = "text/markdown; charset=utf-8"
    elif kind == "txt":
        content = _text_document(answer, metadata, markdown=False).encode("utf-8")
        media_type = "text/plain; charset=utf-8"
    else:
        raise ValueError("Document type must be pdf, docx, txt, md, or json.")
    return GeneratedDocument(
        content=content,
        media_type=media_type,
        filename=f"libra-assist-answer.{kind}",
    )


def _text_document(answer: str, metadata: dict, *, markdown: bool) -> str:
    title = "# Libra Assist answer" if markdown else "LIBRA ASSIST ANSWER"
    details = "\n".join(f"{key}: {value}" for key, value in metadata.items() if value)
    return f"{title}\n\n{details}\n\n{answer}".strip() + "\n"


def _docx(answer: str, metadata: dict) -> bytes:
    from docx import Document

    document = Document()
    document.add_heading("Libra Assist answer", 0)
    for key, value in metadata.items():
        if value:
            paragraph = document.add_paragraph()
            paragraph.add_run(f"{_label(key)}: ").bold = True
            paragraph.add_run(str(value))
    document.add_heading("Answer", level=1)
    for block in re.split(r"\n{2,}", answer):
        document.add_paragraph(block)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def _pdf(answer: str, metadata: dict) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    output = io.BytesIO()
    document = SimpleDocTemplate(
        output,
        pagesize=A4,
        rightMargin=18 * mm,
        leftMargin=18 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="Libra Assist answer",
    )
    styles = getSampleStyleSheet()
    story = [Paragraph("Libra Assist answer", styles["Title"]), Spacer(1, 8)]
    for key, value in metadata.items():
        if value:
            story.append(Paragraph(
                f"<b>{escape(_label(key))}:</b> {escape(str(value))}",
                styles["BodyText"],
            ))
    story.extend([Spacer(1, 12), Paragraph("Answer", styles["Heading1"])])
    for block in re.split(r"\n{2,}", answer):
        safe = escape(block).replace("\n", "<br/>")
        story.extend([Paragraph(safe or " ", styles["BodyText"]), Spacer(1, 6)])
    document.build(story)
    return output.getvalue()


def _label(value: str) -> str:
    return value.replace("_", " ").strip().title()
